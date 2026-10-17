import sys
import os
from pptx import Presentation

def update_deck():
    ppt_path = "ProofStack ppt.pptx"
    if not os.path.exists(ppt_path):
        print(f"Error: {ppt_path} not found.")
        return

    prs = Presentation(ppt_path)
    print(f"Loaded presentation: {len(prs.slides)} slides.")

    # Helper function to find and replace text across all shapes/paragraphs
    def replace_text_in_shapes(shapes, old_text, new_text):
        count = 0
        for s in shapes:
            if s.has_text_frame:
                for p in s.text_frame.paragraphs:
                    if old_text in p.text:
                        p.text = p.text.replace(old_text, new_text)
                        count += 1
            if s.shape_type == 6:  # Group shape
                count += replace_text_in_shapes(s.shapes, old_text, new_text)
        return count

    # Slide 2 Updates (Solution & Prototype)
    slide2 = prs.slides[1]
    replace_text_in_shapes(
        slide2.shapes,
        "Your resume says what you claim. Your repo proves what you did.",
        "Resume says what you claim. Repo proves what you did. ProofStack forensically audits both."
    )
    replace_text_in_shapes(
        slide2.shapes,
        "RECOMMENDATION",
        "COMMIT CONTRACTS & REPO-QUALITY TESTS"
    )
    replace_text_in_shapes(
        slide2.shapes,
        "Private repo skipped · limited analysis",
        "Forensic XYZ Audit: 85% · Repo Quality Tests: Active"
    )

    # Slide 3 Updates (Technical Approach)
    slide3 = prs.slides[2]
    replace_text_in_shapes(
        slide3.shapes,
        "Repo audit, proof map, Agent Trace",
        "Proof map, Google XYZ Studio, Repo Quality Tests, Agent Trace"
    )

    # Slide 4 Updates (Impact & Benefits)
    slide4 = prs.slides[3]
    replace_text_in_shapes(
        slide4.shapes,
        "See exactly what to commit next",
        "Google XYZ bullet scores + exactly what to commit next"
    )
    replace_text_in_shapes(
        slide4.shapes,
        "Verified skills, not keyword resumes",
        "Cryptographically signed code quality attestations vs keyword resumes"
    )

    # Slide 5 Updates (Future Scope & Challenges)
    slide5 = prs.slides[4]
    replace_text_in_shapes(
        slide5.shapes,
        "Evidence of use, not proficiency",
        "Evidence of use vs proficiency (Bridged by Repo Quality Tests)"
    )

    prs.save(ppt_path)
    print("Successfully updated ProofStack ppt.pptx with Deep Forensic & Code Quality Test modules!")

if __name__ == "__main__":
    update_deck()
